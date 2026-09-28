#!/usr/bin/env python3
"""CSV -> a styled .xlsx in the house style (step 2b spreadsheet skill).

Usage: csv_to_xlsx.py <in.csv> <out.xlsx> --tz <IANA zone> [--title "Sheet title"] [--keep-ids]

--tz is the FIRM's timezone (e.g. America/Los_Angeles): timestamps are shown in it,
never in UTC (a 5 pm visit must not read as tomorrow).
"""
import argparse
import csv
import re
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# A money column ENDS in one of these words (total_visits is a count, not money).
MONEY = re.compile(r"(^|_)(cents|amount|balance|price|cost|total)$", re.I)
UUID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$", re.I)
CURRENCY = '"$"#,##0.00'
BAD_TITLE = re.compile(r"[\[\]:*?/\\]")


def label(col: str) -> str:
    col = re.sub(r"_cents$", "", col)
    return " ".join(w.capitalize() for w in col.replace("-", "_").split("_") if w) or col


def as_when(value: str, tz):
    """A date stays a date; a timestamp is converted to the firm's zone. None if neither."""
    if len(value) == 10:
        try:
            return date.fromisoformat(value)
        except ValueError:
            return None
    try:
        d = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    if d.tzinfo is not None:
        d = d.astimezone(tz)
    return d.replace(tzinfo=None)


# Numbers-as-text stay text where a leading zero or a code matters (phones, ZIPs, codes).
TEXTISH = re.compile(r"(phone|zip|postal|code|number|ref)", re.I)
NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def as_number(col: str, value: str):
    if TEXTISH.search(col) or not NUMBER.match(value):
        return None
    return float(value) if "." in value else int(value)


def is_id_column(col: str, rows) -> bool:
    if col == "id" or col.endswith("_id"):
        return True
    values = [r[col] for r in rows if r.get(col)]
    return bool(values) and all(UUID.match(v) for v in values)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--tz", required=True, help="the firm's IANA timezone, e.g. America/Los_Angeles")
    ap.add_argument("--title", default="Export")
    ap.add_argument("--keep-ids", action="store_true")
    a = ap.parse_args()
    tz = ZoneInfo(a.tz)

    with open(a.src, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        cols = list(reader.fieldnames or [])
    if not a.keep_ids:
        cols = [c for c in cols if not is_id_column(c, rows)]

    wb = Workbook()
    ws = wb.active
    ws.title = (BAD_TITLE.sub("-", a.title).strip() or "Export")[:31]
    ws.append([label(c) for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for r in rows:
        out = []
        for c in cols:
            v = r.get(c) or ""
            if v == "":
                out.append(None)
            elif MONEY.search(c):
                try:
                    n = float(v)
                    out.append(n / 100 if c.lower().endswith("cents") else n)
                except ValueError:
                    out.append(v)
            else:
                n = as_number(c, v)
                w = None if n is not None else as_when(v, tz)
                out.append(n if n is not None else w if w is not None else v)
        ws.append(out)

    for i, c in enumerate(cols, start=1):
        letter = get_column_letter(i)
        for cell in ws[letter][1:]:
            if MONEY.search(c) and isinstance(cell.value, (int, float)):
                cell.number_format = CURRENCY
            elif isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
        width = max([len(label(c))] + [len(str(x.value)) for x in ws[letter][1:] if x.value is not None])
        ws.column_dimensions[letter].width = min(max(10, width + 2), 60)

    pulled = datetime.now(tz).strftime("%Y-%m-%d %H:%M")
    ws.append([])
    ws.append([f"Source: {a.title}, {len(rows)} rows, pulled {pulled} ({a.tz})."])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="666666")
    wb.save(a.dst)
    print(f"wrote {a.dst}: {len(rows)} rows, {len(cols)} columns")


if __name__ == "__main__":
    main()
