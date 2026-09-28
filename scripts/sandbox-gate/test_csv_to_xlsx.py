"""Tests for the spreadsheet skill's helper (step 2b). Run: python3 scripts/sandbox-gate/test_csv_to_xlsx.py
Needs openpyxl locally (the sandbox template has it)."""
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime

import openpyxl

SCRIPT = os.path.join(os.path.dirname(__file__), "..", "..", "Box_Setup_Stuff", "seat-glue", "skills", "spreadsheet", "csv_to_xlsx.py")


def build(csv_text, *args):
    d = tempfile.mkdtemp()
    src, dst = os.path.join(d, "in.csv"), os.path.join(d, "out.xlsx")
    with open(src, "w", encoding="utf-8") as f:
        f.write(csv_text)
    subprocess.run([sys.executable, SCRIPT, src, dst, "--tz", "America/Los_Angeles", *args], check=True, capture_output=True)
    ws = openpyxl.load_workbook(dst).active
    return ws, list(ws.iter_rows(values_only=True))


class CsvToXlsx(unittest.TestCase):
    def test_ids_dropped_but_an_all_empty_column_kept(self):
        ws, rows = build("id,customer_id,name,phone\n4f0c1a2e-1111-2222-3333-444455556666,x,Acme,\n")
        self.assertEqual(rows[0], ("Name", "Phone"))

    def test_money_is_a_name_ending_in_a_money_word(self):
        ws, rows = build("balance_cents,total_visits,job_total,cost_code\n18550,3,40,A1\n")
        self.assertEqual(rows[0], ("Balance", "Total Visits", "Job Total", "Cost Code"))
        self.assertEqual(rows[1][:3], (185.5, 3, 40))
        self.assertEqual(ws["A2"].number_format, '"$"#,##0.00')
        self.assertNotEqual(ws["B2"].number_format, '"$"#,##0.00')
        self.assertEqual(rows[1][3], "A1")

    def test_times_are_firm_local_and_dates_stay_dates(self):
        ws, rows = build("starts_at,local_day\n2026-09-28T00:00:00+00:00,2026-09-27\n")
        self.assertEqual(rows[1][0], datetime(2026, 9, 27, 17, 0))  # 5 pm Pacific, not tomorrow
        self.assertEqual(rows[1][1].date() if isinstance(rows[1][1], datetime) else rows[1][1], date(2026, 9, 27))

    def test_counts_become_numbers_but_phones_and_zips_stay_text(self):
        ws, rows = build("visits,phone,postal_code\n3,0412345678,02134\n")
        self.assertEqual(rows[1], (3, "0412345678", "02134"))

    def test_a_title_with_forbidden_characters_is_cleaned(self):
        ws, _ = build("a\n1\n", "--title", "Open jobs 9/27")
        self.assertEqual(ws.title, "Open jobs 9-27")

    def test_header_only_csv_keeps_its_header(self):
        ws, rows = build("name,phone\r\n")
        self.assertEqual(rows[0], ("Name", "Phone"))

    def test_style(self):
        ws, rows = build("name\nZoë Müller\n")
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(rows[-1][0].startswith("Source:"))


if __name__ == "__main__":
    unittest.main()
